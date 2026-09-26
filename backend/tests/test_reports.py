from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.reports import generate_invoices_pdf, generate_invoices_xlsx

TODAY = date(2026, 10, 22)

INVOICES = [
    {
        "invoice_id": "i1", "number": "INV-1001", "client_id": "C01", "status": "open",
        "amount_inr": 50000, "due_inr": 50000, "invoice_date": "2026-09-01", "due_date": "2026-09-10",
    },
    {
        "invoice_id": "i2", "number": "INV-1002", "client_id": "C02", "status": "paid",
        "amount_inr": 20000, "due_inr": 0, "invoice_date": "2026-09-05", "due_date": "2026-09-15",
    },
]
CLIENTS = [
    {"client_id": "C01", "name": "Orion Retail", "tier": "Regular"},
    {"client_id": "C02", "name": "Saffron Threads", "tier": "VIP"},
]


def test_xlsx_has_invoices_and_summary_sheets():
    content = generate_invoices_xlsx(INVOICES, CLIENTS, today=TODAY)
    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == ["Invoices", "Summary"]


def test_xlsx_computes_days_overdue_from_the_given_today_not_real_time():
    content = generate_invoices_xlsx(INVOICES, CLIENTS, today=TODAY)
    wb = load_workbook(BytesIO(content))
    ws = wb["Invoices"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    by_number = {r[0]: r for r in rows}
    assert by_number["INV-1001"][-1] == 42  # 2026-09-10 -> 2026-10-22


def test_xlsx_paid_invoice_has_zero_days_overdue_regardless_of_due_date():
    content = generate_invoices_xlsx(INVOICES, CLIENTS, today=TODAY)
    wb = load_workbook(BytesIO(content))
    ws = wb["Invoices"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    by_number = {r[0]: r for r in rows}
    assert by_number["INV-1002"][-1] == 0


def test_xlsx_resolves_client_name_and_tier():
    content = generate_invoices_xlsx(INVOICES, CLIENTS, today=TODAY)
    wb = load_workbook(BytesIO(content))
    ws = wb["Invoices"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    by_number = {r[0]: r for r in rows}
    assert by_number["INV-1001"][1] == "Orion Retail"
    assert by_number["INV-1001"][2] == "Regular"


def test_xlsx_summary_totals_match_invoice_amounts():
    content = generate_invoices_xlsx(INVOICES, CLIENTS, today=TODAY)
    wb = load_workbook(BytesIO(content))
    summary = wb["Summary"]
    values = {row[0]: row[1] for row in summary.iter_rows(min_row=2, values_only=True)}
    assert values["Total invoices"] == 2
    assert values["Total invoiced (INR)"] == 70000
    assert values["Total outstanding (INR)"] == 50000


def test_pdf_starts_with_pdf_magic_bytes_and_is_non_trivial():
    content = generate_invoices_pdf(INVOICES, CLIENTS, today=TODAY)
    assert content[:5] == b"%PDF-"
    assert len(content) > 500


def test_empty_invoice_list_still_produces_valid_files():
    xlsx = generate_invoices_xlsx([], [], today=TODAY)
    wb = load_workbook(BytesIO(xlsx))
    assert wb.sheetnames == ["Invoices", "Summary"]

    pdf = generate_invoices_pdf([], [], today=TODAY)
    assert pdf[:5] == b"%PDF-"
