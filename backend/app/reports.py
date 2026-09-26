"""Generates a downloadable Excel/PDF payments report from current invoice data - for
sending to an accountant/CA, not for the agent's own use. Pure formatting over data
already in our db; no external API calls, so this never goes through swy/executor.py.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.money import format_inr


# reportlab's default fonts (Helvetica etc.) have no glyph for "₹" - it renders as a
# solid black box in the PDF. The dashboard can use format_inr()'s "₹" directly since
# browsers/HTML have full Unicode font fallback, but the PDF needs a plain-ASCII prefix.
def _format_inr_pdf(amount: int) -> str:
    return format_inr(amount).replace("₹", "Rs. ")

COLUMNS = [
    "Invoice", "Client", "Tier", "Invoice Date", "Due Date",
    "Amount (INR)", "Paid (INR)", "Due (INR)", "Status", "Days Overdue",
]


def _rows(invoices: list[dict], clients: list[dict], *, today: date) -> list[list]:
    client_names = {c["client_id"]: c["name"] for c in clients}
    client_tiers = {c["client_id"]: c["tier"] for c in clients}

    rows = []
    for inv in invoices:
        due_date = date.fromisoformat(inv["due_date"])
        days_overdue = max(0, (today - due_date).days) if inv["status"] != "paid" else 0
        paid_inr = inv["amount_inr"] - inv["due_inr"]
        rows.append([
            inv["number"],
            client_names.get(inv["client_id"], inv["client_id"]),
            client_tiers.get(inv["client_id"], ""),
            inv["invoice_date"],
            inv["due_date"],
            inv["amount_inr"],
            paid_inr,
            inv["due_inr"],
            inv["status"].upper(),
            days_overdue,
        ])
    rows.sort(key=lambda r: r[-1], reverse=True)
    return rows


def generate_invoices_xlsx(invoices: list[dict], clients: list[dict], *, today: date) -> bytes:
    rows = _rows(invoices, clients, today=today)

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"

    ws.append([f"MunimJi - Payments Report - {today.isoformat()}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(COLUMNS)
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append(row)

    money_cols = {6, 7, 8}  # Amount, Paid, Due (1-indexed)
    for row_idx in range(header_row + 1, header_row + 1 + len(rows)):
        for col_idx in money_cols:
            ws.cell(row=row_idx, column=col_idx).number_format = "#,##0"

    for col_idx, header in enumerate(COLUMNS, start=1):
        width = max(len(header), 12) + 2
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    total_amount = sum(r[5] for r in rows)
    total_due = sum(r[7] for r in rows)
    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary["A1"].font = Font(bold=True)
    summary["B1"].font = Font(bold=True)
    summary.append(["Report date", today.isoformat()])
    summary.append(["Total invoices", len(rows)])
    summary.append(["Total invoiced (INR)", total_amount])
    summary.append(["Total outstanding (INR)", total_due])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 20

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_invoices_pdf(invoices: list[dict], clients: list[dict], *, today: date) -> bytes:
    rows = _rows(invoices, clients, today=today)
    total_amount = sum(r[5] for r in rows)
    total_due = sum(r[7] for r in rows)

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Heading1"], fontSize=16, spaceAfter=4)
    meta_style = ParagraphStyle("Meta", parent=styles["Normal"], textColor=colors.grey)

    elements = [
        Paragraph("MunimJi - Payments Report", title_style),
        Paragraph(f"Generated {today.isoformat()} - Kaarigar Studio", meta_style),
        Spacer(1, 4 * mm),
        Paragraph(
            f"Total invoiced: {_format_inr_pdf(total_amount)}  |  Total outstanding: {_format_inr_pdf(total_due)}"
            f"  |  Invoices: {len(rows)}",
            styles["Normal"],
        ),
        Spacer(1, 6 * mm),
    ]

    table_data = [COLUMNS] + [
        [
            r[0], r[1], r[2], r[3], r[4],
            _format_inr_pdf(r[5]), _format_inr_pdf(r[6]), _format_inr_pdf(r[7]), r[8], str(r[9]),
        ]
        for r in rows
    ]
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (5, 1), (9, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f4f8")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(table)

    doc.build(elements)
    return buf.getvalue()
